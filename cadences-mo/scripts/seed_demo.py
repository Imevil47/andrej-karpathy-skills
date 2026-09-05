# -*- coding: utf-8 -*-
"""Seed the validation scenarios A to E into a copy of the v2.0 workbook."""
import datetime as dt
import sys

import openpyxl

from build import (LIST_FIRST, STD_FIRST, EMP_FIRST, RUNS_FIRST, BASE_FIRST,
                   TEST_FIRST, ARRET_FIRST)

NOMS = [
    "AMRANI FATIMA", "BENALI KHADIJA", "CHAKIR SALMA", "DAOUDI NAIMA",
    "EL IDRISSI ZINEB", "FAHMI RACHIDA", "GHANEM SOUAD", "HAMDI MERIEM",
    "IDRISSI LATIFA", "JAMAL HANANE", "KABBAJ SANAA", "LAMRANI NADIA",
    "MOUJAHID ASMAA", "NACIRI IMANE", "OUAZZANI HAFIDA", "QUADIRI SIHAM",
    "RAMI BOUCHRA", "SAIDI KAOUTAR", "TAZI MALIKA", "URRIOL SAADIA",
    "VADEL LOUBNA", "WAHBI SAMIRA", "YOUSFI HANAA", "ZAHIDI RAJAA",
    "ABBOU LEILA", "BOUZIDI SAFAA", "CHERKAOUI AMAL", "DRISSI WAFAA",
    "ESSADIK HOUDA", "FARSI NOUHAILA", "GUESSOUS SOUMIA", "HILALI JAMILA",
    "IRAQI YASMINE", "JEBBOUR HAYAT", "KHALIL FADWA", "LOUKILI DOUNIA",
    "MANSOURI SOUKAINA", "NASSIRI GHIZLANE", "OMARI CHAIMAE", "PAKI MOUNIA",
]
MATRICULES = ["M%03d" % (i + 1) for i in range(len(NOMS))]


def seed_parametres(ws):
    for i, v in enumerate(["MARQUE A", "MARQUE B"]):
        ws.cell(LIST_FIRST + i, 4, v)                  # MARQUE
    for i, v in enumerate(["1/4 CLUB", "1/5 OVALE"]):
        ws.cell(LIST_FIRST + i, 5, v)                  # FORMAT
    for i, v in enumerate(["MOULE 1", "MOULE 2"]):
        ws.cell(LIST_FIRST + i, 7, v)                  # MOULE
    standards = [
        ("SARDINE", "FMHOEV BIO", "MARQUE A", "1/4 CLUB", "HGT", "", "GRATTAGE + REMPLISSAGE", 112, "OUI"),
        ("SARDINE", "FMHOEV BIO", "MARQUE A", "1/4 CLUB", "HG", "", "GRATTAGE + REMPLISSAGE", 105, "OUI"),
        ("MAQUEREAU", "SPSA HO", "MARQUE B", "1/5 OVALE", "", "MOULE 1", "GRATTAGE", 90, "OUI"),
        ("MAQUEREAU", "SPSA HO", "MARQUE B", "1/5 OVALE", "", "MOULE 1", "REMPLISSAGE", 130, "OUI"),
    ]
    for i, row in enumerate(standards):
        for j, v in enumerate(row):
            if v != "":
                ws.cell(STD_FIRST + 1 + i, 1 + j, v)
    for i, (mat, nom) in enumerate(zip(MATRICULES, NOMS)):
        ws.cell(EMP_FIRST + i, 1, mat)
        ws.cell(EMP_FIRST + i, 2, nom)
        ws.cell(EMP_FIRST + i, 3, "OUI")


RUN_A = "RUN-20260904-01"
RUN_B = "RUN-20260905-01"


def seed_runs(ws):
    runs = [
        (RUN_A, dt.date(2026, 9, 4), dt.time(8, 0), dt.time(16, 0), "SARDINE",
         "FMHOEV BIO", "MARQUE A", "1/4 CLUB", "LOT-2609-A", "HGT", "",
         ["GRATTAGE + REMPLISSAGE"] * 3 + ["INACTIVE"] * 5, "ACTIF"),
        (RUN_B, dt.date(2026, 9, 5), dt.time(8, 0), dt.time(16, 0), "MAQUEREAU",
         "SPSA HO", "MARQUE B", "1/5 OVALE", "LOT-2609-B", "", "MOULE 1",
         ["GRATTAGE", "GRATTAGE", "REMPLISSAGE", "REMPLISSAGE"] + ["INACTIVE"] * 4,
         "ACTIF"),
    ]
    for i, run in enumerate(runs):
        r = RUNS_FIRST + i
        flat = list(run[:11]) + run[11] + [run[12]]
        for j, v in enumerate(flat):
            if v != "":
                ws.cell(r, 1 + j, v)


def seed_base(ws):
    """Rows mimic exactly what SAISIE CONTROLE pastes: columns B..I only."""
    rows = []

    def add(date, heure, run, tour, ligne, mats, boxes, effectif):
        for mat, nb in zip(mats, boxes):
            rows.append([date, heure, run, tour, ligne, mat, nb, effectif])

    d1 = dt.date(2026, 9, 4)
    l1 = MATRICULES[0:12]
    l2 = MATRICULES[12:24]
    l3 = MATRICULES[24:36]
    # Scenario A - sardine, three lines, three control rounds.
    add(d1, dt.time(9, 5), RUN_A, "TOUR 01", "L1", l1, [0] * 12, 12)
    add(d1, dt.time(9, 5), RUN_A, "TOUR 01", "L2", l2, [0] * 12, 12)
    add(d1, dt.time(9, 5), RUN_A, "TOUR 01", "L3", l3, [0] * 12, 12)
    t2 = [112, 110, 105, 100, 95, 90, 115, 108, 102, 99, 88, 120]
    add(d1, dt.time(10, 3), RUN_A, "TOUR 02", "L1", l1, t2, 12)
    # Scenario C - the same matricule entered twice in the same RUN+TOUR+LIGNE.
    rows.append([d1, dt.time(10, 3), RUN_A, "TOUR 02", "L1", "M005", 95, 12])
    # Scenario D - incomplete line: 10 of 12 women controlled.
    add(d1, dt.time(10, 3), RUN_A, "TOUR 02", "L2",
        l2[:10], [104, 118, 96, 111, 107, 92, 113, 101, 109, 98], 12)
    add(d1, dt.time(10, 3), RUN_A, "TOUR 02", "L3", l3,
        [106, 114, 99, 103, 117, 94, 108, 111, 97, 105, 119, 101], 12)
    add(d1, dt.time(11, 0), RUN_A, "TOUR 03", "L1", l1,
        [108, 112, 101, 97, 104, 93, 118, 110, 99, 106, 91, 115], 12)

    # Scenario B - maquereau, grattage and remplissage split across lines.
    d2 = dt.date(2026, 9, 5)
    g1, g2 = MATRICULES[0:6], MATRICULES[6:12]
    r1, r2 = MATRICULES[12:18], MATRICULES[18:24]
    for ligne, mats in (("L1", g1), ("L2", g2), ("L3", r1), ("L4", r2)):
        add(d2, dt.time(8, 30), RUN_B, "TOUR 01", ligne, mats, [0] * 6, 6)
    add(d2, dt.time(9, 32), RUN_B, "TOUR 02", "L1", g1, [95, 88, 92, 84, 97, 90], 6)
    add(d2, dt.time(9, 32), RUN_B, "TOUR 02", "L2", g2, [86, 91, 79, 94, 88, 93], 6)
    add(d2, dt.time(9, 32), RUN_B, "TOUR 02", "L3", r1, [136, 128, 141, 122, 133, 130], 6)
    add(d2, dt.time(9, 32), RUN_B, "TOUR 02", "L4", r2, [125, 138, 119, 132, 127, 135], 6)

    for i, row in enumerate(rows):
        r = BASE_FIRST + i
        for j, v in enumerate(row):
            ws.cell(r, 2 + j, v)
    return len(rows)


def seed_others(wb):
    ws = wb["TEST RENDEMENT"]
    tests = [
        (dt.date(2026, 9, 4), "M001", "SARDINE", "FMHOEV BIO", "MARQUE A",
         "1/4 CLUB", "HGT", "", 10.0, 7.8, "TEST INDIVIDUEL",
         "Scenario E — rendement attendu 78,0 %"),
        (dt.date(2026, 9, 5), "M007", "MAQUEREAU", "SPSA HO", "MARQUE B",
         "1/5 OVALE", "", "MOULE 1", 12.5, 9.375, "TEST COMPARATIF", ""),
    ]
    for i, t in enumerate(tests):
        r = TEST_FIRST + i
        ws.cell(r, 2, t[0]); ws.cell(r, 3, t[1])
        for j, v in enumerate(t[2:8]):
            if v != "":
                ws.cell(r, 5 + j, v)
        ws.cell(r, 11, t[8]); ws.cell(r, 12, t[9])
        ws.cell(r, 14, t[10]); ws.cell(r, 15, t[11])

    ws = wb["ARRETS"]
    arrets = [
        (dt.date(2026, 9, 4), RUN_A, "L1", dt.time(12, 0), dt.time(12, 30), "PAUSE", ""),
        (dt.date(2026, 9, 4), RUN_A, "L2", dt.time(10, 15), dt.time(10, 48), "PANNE",
         "Sertisseuse — intervention maintenance"),
        (dt.date(2026, 9, 5), RUN_B, "L3", dt.time(9, 5), dt.time(9, 20),
         "MANQUE MATIERE", ""),
    ]
    for i, a in enumerate(arrets):
        r = ARRET_FIRST + i
        for j, v in enumerate(a):
            if v != "":
                ws.cell(r, 1 + j, v)

    ws = wb["SAISIE CONTROLE"]
    ws["C4"] = RUN_A
    ws["C5"] = "TOUR 04"
    ws["C6"] = "L1"
    ws["C7"] = dt.time(12, 0)
    ws["C8"] = 12
    for i, mat in enumerate(MATRICULES[0:5]):
        ws.cell(13 + i, 2, mat)
        ws.cell(13 + i, 4, 100 + i)
    ws["B18"] = "M003"          # duplicate inside the entry grid
    ws["D18"] = 90

    ws = wb["SYNTHESE CONTROLES"]
    ws["C4"] = RUN_A
    ws["C5"] = "TOUR 02"

    ws = wb["RECHERCHE MATRICULE"]
    ws["B4"] = "M001"


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    seed_parametres(wb["PARAMETRES"])
    seed_runs(wb["RUNS"])
    n = seed_base(wb["BASE CONTROLES"])
    seed_others(wb)
    wb.save(dst)
    print("seeded %d controles -> %s" % (n, dst))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
