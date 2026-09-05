"""Shared visual identity, inherited from Cadences MO Grattage v1.3."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

NAVY = "0A0F1A"
ORANGE = "FF7F3F"
TEAL = "17A598"
CREAM = "FFF6E5"
LIGHT = "F2F5F7"
WHITE = "FFFFFF"
BLUE_TXT = "0000FF"
GREY_TXT = "5A6672"

OK_FILL, OK_TXT = "D6F2EC", "0E6B62"
WARN_FILL, WARN_TXT = "FFEBDA", "B04E14"
BAD_FILL, BAD_TXT = "FBD9D9", "9E1C1C"
NEUTRAL_FILL, NEUTRAL_TXT = "F2F5F7", "5A6672"

FMT_DATE = "dd/mm/yyyy"
FMT_TIME = "hh:mm"
FMT_INT = "#,##0"
FMT_DEC = "#,##0.0"
FMT_PCT = "0.0%"
FMT_KG = "#,##0.000"
FMT_GAP = "\\+#,##0.0;\\-#,##0.0;0.0"

_THIN = Side(style="thin", color="BFC7CF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def title(ws, ref, text):
    """Workbook-level banner across `ref` (e.g. 'A1:H1')."""
    ws.merge_cells(ref)
    c = ws[ref.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[c.row].height = 30
    for cell in ws[ref.replace(":", ":")][0]:
        cell.fill = _fill(NAVY)


def section(ws, ref, text):
    ws.merge_cells(ref)
    c = ws[ref.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT_NAME, size=11, bold=True, color=WHITE)
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[c.row].height = 21
    for cell in ws[ref][0]:
        cell.fill = _fill(NAVY)


def note(ws, ref, text):
    ws.merge_cells(ref)
    c = ws[ref.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT_NAME, size=9, italic=True, color=GREY_TXT)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def headers(ws, row, first_col, labels, colour=ORANGE):
    for i, label in enumerate(labels):
        c = ws.cell(row, first_col + i, label)
        c.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
        c.fill = _fill(colour)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 33


def input_cell(ws, coord, number_format, align="center"):
    c = ws[coord]
    c.font = Font(name=FONT_NAME, size=10, color=BLUE_TXT)
    c.fill = _fill(CREAM)
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.number_format = number_format
    return c


def calc_cell(ws, coord, number_format, align="center", bold=False):
    c = ws[coord]
    c.font = Font(name=FONT_NAME, size=10, bold=bold, color=NAVY)
    c.fill = _fill(WHITE)
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.number_format = number_format
    return c


def label_cell(ws, coord, text, align="right"):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
    c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def band_cell(ws, coord, number_format, align="center", bold=True):
    c = ws[coord]
    c.font = Font(name=FONT_NAME, size=10, bold=bold, color=NAVY)
    c.fill = _fill(LIGHT)
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.number_format = number_format
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def rows_height(ws, first, last, height):
    for r in range(first, last + 1):
        ws.row_dimensions[r].height = height
